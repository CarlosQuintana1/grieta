"""Arma el Excel para revisar el modelo: los registros y lo predicho contra lo real (80/20 por tiempo).

    python exportar_excel.py                     # lee datos/dataset.csv
    python exportar_excel.py --salida otro.xlsx
"""

import argparse
import json
import pathlib
import re
import zipfile
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from sklearn.linear_model import Ridge

from entrenar_modelo import MODELO, PLATAFORMAS, cargar, variables

APPS = {"uber": "Uber", "didi": "DiDi", "indrive": "inDrive"}
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
PRUEBA = 0.2
LETRA = Font(name="Arial", size=10)
NEGRITA = Font(name="Arial", size=10, bold=True)
TITULO = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="1F3A5F")
PESOS = "$#,##0.00"

# (encabezado, ancho, formato, nota)
REGISTROS = [
    ("Registro", 10, None, "Un registro es la misma ruta cotizada en 2 o 3 apps al mismo tiempo. "
                           "Cada app va en su propia fila."),
    ("Fecha", 12, "dd/mm/yyyy", None),
    ("Hora", 8, "hh:mm", "Hora local de Querétaro."),
    ("Día", 11, None, None),
    ("Franja", 8, None, "Bloque de horas, en hora local."),
    ("Fin de semana", 9, None, None),
    ("Festivo", 9, None, None),
    ("Origen", 34, None, None),
    ("Destino", 34, None, None),
    ("Ruta", 7, None, "Número de la ruta en el catálogo: cada sentido es una ruta distinta."),
    ("Km", 7, "0.0", "Distancia de la ruta según OSRM."),
    ("Minutos sin tráfico", 10, "0.0", "Duración de la ruta sin tráfico según OSRM."),
    ("Zona origen", 17, None, "Hexágono H3 de resolución 7, de unos 5 km²."),
    ("Zona destino", 17, None, "Hexágono H3 de resolución 7, de unos 5 km²."),
    ("Temperatura (°C)", 12, "0.0", "Histórico por hora en el lugar de origen, de Open-Meteo."),
    ("Lluvia (mm)", 9, "0.0", "Histórico por hora en el lugar de origen, de Open-Meteo."),
    ("App", 9, None, None),
    ("Servicio", 18, None, None),
    ("Precio real ($)", 12, PESOS, "Lo que mostró la app al cotizar, capturado a mano."),
    ("Piloto", 8, None, "Capturas de antes de fijar los puntos de referencia (23 de septiembre): "
                        "sus km pueden no coincidir del todo."),
    ("Conjunto", 14, None, "80/20 por tiempo: el 80% más antiguo de los registros entrena el modelo "
                           "y el 20% más reciente se usa para probarlo."),
]

PREDICHO = [
    ("Registro", 10, None, None),
    ("Fecha", 12, "dd/mm/yyyy", None),
    ("Hora", 8, "hh:mm", None),
    ("Origen", 34, None, None),
    ("Destino", 34, None, None),
    ("App", 9, None, None),
    ("Precio real ($)", 12, PESOS, None),
    ("Precio predicho ($)", 12, PESOS, "Regresión lineal (Ridge, alfa 1), una por app, con la ruta y la "
                                       "franja horaria como variables. Se entrenó solo con los registros "
                                       "marcados como Entrenamiento en la hoja Registros."),
    ("Diferencia ($)", 12, PESOS, "Predicho menos real: positivo si el modelo se pasó, negativo si se "
                                  "quedó corto."),
    ("Error absoluto ($)", 12, PESOS, None),
    ("Error (%)", 9, "0.0%", "Error absoluto entre el precio real."),
]


def nombres_de_ruta() -> dict[int, tuple[str, str]]:
    catalogo = MODELO.parent / "captura" / "data" / "lugares.json"
    lugares = {l["id"]: l["nombre"] for l in json.loads(catalogo.read_text())["lugares"]}
    rutas = json.loads((MODELO / "datos" / "rutas_medidas.json").read_text())["rutas"]
    return {r["ruta_id"]: (lugares[r["origen"]], lugares[r["destino"]]) for r in rutas}


def separar_y_predecir(df: pd.DataFrame) -> pd.DataFrame:
    # Se prueba con lo mas reciente, como se usaria el modelo: predecir lo que todavia no pasa.
    orden = df.drop_duplicates("momento_id").sort_values("timestamp")["momento_id"]
    numero = {m: i + 1 for i, m in enumerate(orden)}
    corte = round(len(orden) * (1 - PRUEBA))
    df = df.assign(registro=df["momento_id"].map(numero))
    df["conjunto"] = ["Prueba" if n > corte else "Entrenamiento" for n in df["registro"]]
    X = variables(df)
    df["predicho"] = float("nan")
    for p in PLATAFORMAS:
        es_p = df["plataforma"] == p
        entrena = es_p & (df["conjunto"] == "Entrenamiento")
        prueba = es_p & (df["conjunto"] == "Prueba")
        if entrena.sum() < 5 or not prueba.any():
            continue
        modelo = Ridge(alpha=1.0).fit(X[entrena], df.loc[entrena, "precio_mxn"])
        df.loc[prueba, "predicho"] = modelo.predict(X[prueba]).round(2)
    df["orden_app"] = df["plataforma"].map({p: i for i, p in enumerate(APPS)})
    return df.sort_values(["registro", "orden_app"])


def encabezar(ws, columnas):
    for i, (titulo, ancho, _, nota) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i, value=titulo)
        celda.font, celda.fill = TITULO, FONDO
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        if nota:
            celda.comment = Comment(nota, "Equipo")
        ws.column_dimensions[celda.column_letter].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def escribir(ws, columnas, filas):
    for r, fila in enumerate(filas, start=2):
        for c, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=valor)
            celda.font = LETRA
            if columnas[c - 1][2]:
                celda.number_format = columnas[c - 1][2]
    ultima = ws.cell(row=1, column=len(columnas)).column_letter
    ws.auto_filter.ref = f"A1:{ultima}{len(filas) + 1}"


def vacio(valor):
    return None if pd.isna(valor) else valor


def guardar(wb: Workbook, salida: pathlib.Path, hoja: str, calculados: dict[str, float]):
    # Excel recalcula al abrir, pero las vistas previas no: cada formula se guarda con su resultado.
    wb.save(salida)
    nombre = f"xl/worksheets/sheet{wb.sheetnames.index(hoja) + 1}.xml"
    with zipfile.ZipFile(salida) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    partes[nombre] = re.sub(
        r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v ?/></c>',
        lambda m: f'<c r="{m[1]}"{m[2]}><f>{m[3]}</f><v>{calculados[m[1]]}</v></c>'
        if m[1] in calculados else m[0],
        partes[nombre].decode(),
    ).encode()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for n, datos in partes.items():
            z.writestr(n, datos)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--entrada", default=str(MODELO / "datos" / "dataset.csv"))
    ap.add_argument("--salida", default=str(MODELO / "datos" / f"tarifas-modelo-{date.today()}.xlsx"))
    args = ap.parse_args()

    df = separar_y_predecir(cargar(pathlib.Path(args.entrada)))
    hora = pd.to_datetime(df["timestamp"], utc=True, format="ISO8601").dt.tz_convert("America/Mexico_City")
    df["fecha"], df["hora"] = hora.dt.date, hora.dt.floor("min").dt.time
    nombres = nombres_de_ruta()
    si_no = lambda v: "Sí" if str(v).lower() == "true" else "No"

    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    registros = wb.active
    registros.title = "Registros"
    encabezar(registros, REGISTROS)
    escribir(registros, REGISTROS, [
        [f.registro, f.fecha, f.hora, DIAS[f.dia_semana], str(f.franja), si_no(f.es_findesemana),
         si_no(f.es_festivo), *nombres[f.ruta_id], f.ruta_id, f.distancia_km, f.duracion_libre_min,
         f.h3_o, f.h3_d, vacio(f.temp_c), vacio(f.lluvia_mm), APPS[f.plataforma], f.tipo_vehiculo,
         f.precio_mxn, "Sí" if f.fuente == "piloto" else "No", f.conjunto]
        for f in df.itertuples()
    ])

    predicho = wb.create_sheet("Predicho vs real")
    encabezar(predicho, PREDICHO)
    prueba = df[df["predicho"].notna()]
    escribir(predicho, PREDICHO, [
        [f.registro, f.fecha, f.hora, *nombres[f.ruta_id], APPS[f.plataforma], f.precio_mxn, f.predicho,
         f"=H{r}-G{r}", f"=ABS(I{r})", f'=IF(G{r}>0,J{r}/G{r},"")']
        for r, f in enumerate(prueba.itertuples(), start=2)
    ])

    n = len(prueba) + 1
    resumen = [("App", "Error promedio ($)", "Filas")] + [
        (app, f'=IFERROR(AVERAGEIFS($J$2:$J${n},$F$2:$F${n},M{r}),"")', f"=COUNTIF($F$2:$F${n},M{r})")
        for r, app in enumerate(APPS.values(), start=2)
    ] + [("Todas", f"=AVERAGE($J$2:$J${n})", f"=COUNTA($F$2:$F${n})")]
    for r, fila in enumerate(resumen, start=1):
        for c, valor in enumerate(fila, start=13):
            celda = predicho.cell(row=r, column=c, value=valor)
            celda.font = TITULO if r == 1 else (NEGRITA if c == 13 else LETRA)
            if r == 1:
                celda.fill = FONDO
            elif c == 14:
                celda.number_format = PESOS
    predicho["N1"].comment = Comment("Promedio del error absoluto en la prueba: el 20% más reciente "
                                     "de los registros, que el modelo no vio al entrenar.", "Equipo")
    for letra, ancho in (("L", 3), ("M", 10), ("N", 18), ("O", 8)):
        predicho.column_dimensions[letra].width = ancho

    calculados = {}
    for r, f in enumerate(prueba.itertuples(), start=2):
        diferencia = f.predicho - f.precio_mxn
        calculados.update({f"I{r}": round(diferencia, 6), f"J{r}": round(abs(diferencia), 6),
                           f"K{r}": round(abs(diferencia) / f.precio_mxn, 6)})
    errores = (prueba["predicho"] - prueba["precio_mxn"]).abs()
    for r, p in enumerate(APPS, start=2):
        de_p = errores[prueba["plataforma"] == p]
        calculados[f"O{r}"] = len(de_p)
        if not de_p.empty:
            calculados[f"N{r}"] = round(de_p.mean(), 6)
    calculados.update({"N5": round(errores.mean(), 6), "O5": len(prueba)})

    salida = pathlib.Path(args.salida)
    guardar(wb, salida, "Predicho vs real", calculados)
    print(f"{df['registro'].nunique()} registros ({len(df)} filas); prueba: {prueba['registro'].nunique()} "
          f"registros ({len(prueba)} filas) -> {salida.name}")
    for r, app in enumerate(APPS.values(), start=2):
        if f"N{r}" in calculados:
            print(f"  {app:8s} error promedio ${calculados[f'N{r}']:.2f}  ({calculados[f'O{r}']} filas)")


if __name__ == "__main__":
    main()
